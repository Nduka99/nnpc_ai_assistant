"""Spreadsheets read into elements in the shape PDF tables take. Workbooks invented."""

from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
import pytest
import xlrd

from petroknowledge import sheets


def test_cells_read_as_written_without_float_noise() -> None:
    assert sheets.text(None) == ""
    assert sheets.text(datetime(2022, 7, 17)) == "2022-07-17"
    assert sheets.text(datetime(2022, 7, 17, 9, 30)) == "2022-07-17T09:30:00"
    assert sheets.text(date(2022, 7, 17)) == "2022-07-17"
    assert sheets.text(True) == "True"
    assert sheets.text(5425548671.639999) == "5425548671.64"
    assert sheets.text(2285.24964017758) == "2285.24964017758"
    assert sheets.text(7) == "7"
    assert sheets.text("  GAS \n FLARED ") == "GAS FLARED"


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    """A gas-data sheet (merged title, a two-level header with units, a blank
    margin column), an empty sheet, a text-only table and a sheet of notes."""
    book = openpyxl.Workbook()
    gas = book.active
    assert gas is not None
    gas.title = "Gas Data"
    rows: list[list[Any]] = [
        [None, "GAS PRODUCTION DATA", None, None, None],
        [None, "AUGUST, 2022"],
        [None, "COMPANY", "GAS PRODUCED", None, "REMARKS"],
        [None, None, "AG", "NAG", None],
        [None, None, "(MMSCF)", "(MMSCF)", None],
        [None, "SHELL", 2285.24964017758, 19298.3, "ok"],
        [],
        [None, "MOBIL", 25984.747, 0, None],
    ]
    for row in rows:
        gas.append(row)
    gas.merge_cells("B1:E1")  # the title across the table
    gas.merge_cells("C3:D3")  # a group label over two columns
    book.create_sheet("Empty")
    partners = book.create_sheet("Partners")
    for row in [["Partner", "Category"], ["Duke Oil", "A"], ["Vitol", "A"]]:
        partners.append(row)
    notes = book.create_sheet("Notes")
    for line in ["SELECTION PROCESS", "1. Tender published", "2. Bids opened"]:
        notes.append([line])
    path = tmp_path / ("a" * 64)  # stored files are named by SHA-256, no extension
    book.save(path)
    return path


def test_a_workbook_reads_as_titles_tables_headers_and_rows(workbook: Path) -> None:
    found = sheets.read_book(workbook)

    gas = [element for element in found if element.page == 1]
    caption = "Gas Data — GAS PRODUCTION DATA — AUGUST, 2022"
    assert [(e.kind, e.text) for e in gas[:2]] == [
        ("heading", "GAS PRODUCTION DATA"),
        ("paragraph", "AUGUST, 2022"),
    ]
    assert gas[1].section == ("Gas Data", "GAS PRODUCTION DATA")
    header = "COMPANY | GAS PRODUCED AG (MMSCF) | GAS PRODUCED NAG (MMSCF) | REMARKS"
    assert gas[2].kind == "table" and gas[2].text.startswith(f"{caption}\n{header}")
    assert gas[2].grid[1] == ("SHELL", "2285.24964017758", "19298.3", "ok")
    assert (gas[3].kind, gas[3].text) == ("table_header", header)
    assert [e.text for e in gas[4:]] == [
        f"{caption} › SHELL — GAS PRODUCED AG (MMSCF): 2285.24964017758; "
        "GAS PRODUCED NAG (MMSCF): 19298.3; REMARKS: ok",
        f"{caption} › MOBIL — GAS PRODUCED AG (MMSCF): 25984.747; "
        "GAS PRODUCED NAG (MMSCF): 0",
    ]
    assert gas[5].row == 1 and gas[5].table == 1
    assert [e.page for e in found if e.page != 1] == [3, 3, 3, 3, 4, 4, 4]  # no sheet 2
    partners = [e.text for e in found if e.page == 3]
    assert partners[1:] == [
        "Partner | Category",
        "Partners › Duke Oil — Category: A",
        "Partners › Vitol — Category: A",
    ]
    notes = [(e.kind, e.text) for e in found if e.page == 4]
    assert notes == [
        ("heading", "SELECTION PROCESS"),
        ("paragraph", "1. Tender published"),
        ("paragraph", "2. Bids opened"),
    ]


class Cell:
    def __init__(self, ctype: int, value: Any) -> None:
        self.ctype = ctype
        self.value = value


class Sheet:
    """An old-format sheet: a title merged across both columns, then a table whose
    second column holds an Excel date serial."""

    name = "Liftings"
    merged_cells = [(0, 1, 0, 2)]  # rows 0-1, columns 0-2 (end exclusive)
    rows = [
        [Cell(xlrd.XL_CELL_TEXT, "LIFTINGS"), Cell(xlrd.XL_CELL_EMPTY, "")],
        [Cell(xlrd.XL_CELL_TEXT, "Customer"), Cell(xlrd.XL_CELL_TEXT, "BL date")],
        [Cell(xlrd.XL_CELL_TEXT, "Duke Oil"), Cell(xlrd.XL_CELL_DATE, 44774.0)],
    ]
    nrows = 3

    def row(self, index: int) -> list[Cell]:
        return self.rows[index]


class Book:
    datemode = 0

    def sheets(self) -> list[Sheet]:
        return [Sheet()]


def test_an_old_xls_workbook_is_read_with_merges_and_dates(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "old.xls"
    path.write_bytes(sheets.OLE2 + bytes(60))
    opened: list[tuple[Path, bool]] = []

    def open_workbook(name: Path, formatting_info: bool) -> Book:
        opened.append((name, formatting_info))
        return Book()

    monkeypatch.setattr(xlrd, "open_workbook", open_workbook)

    found = sheets.read_book(path)

    assert opened == [(path, True)]  # merged cells need the formatting information
    assert [(e.kind, e.text) for e in found] == [
        ("heading", "LIFTINGS"),
        ("table", "Liftings — LIFTINGS\nCustomer | BL date\nDuke Oil | 2022-08-01"),
        ("table_header", "Customer | BL date"),
        ("table_row", "Liftings — LIFTINGS › Duke Oil — BL date: 2022-08-01"),
    ]


def test_header_blocks_are_found_wherever_they_stand() -> None:
    """A row of column numbers before the header, a second block lower down beside
    a side table, and text rows that no figure follows (rows, not a header)."""
    rows: list[list[Any]] = [
        ["ROYALTY LIFTINGS FOR JULY 2022"],
        [1, "2", "3", None, None, None],  # column numbers, not a header
        ["CUSTOMER", "VESSEL", "QTY IN BBLS", None, None, None],
        [None, None, "(BBLS)", None, None, None],  # its units, under one column
        ["STERLING", "STATIA", 200000, None, None, None],
        ["RENTAL (NAIRA)"],  # a section's name: one text, a row
        ["CUSTOMER", "VESSEL", "EXCH RATE", None, "COMPANIES", "COMPANY LIFTING"],
        [None, None, None, None, None, None],
        ["OANDO", "DELOS", 401.3, None, "NPDC/NACONDE", 0],
        ["Prepared by", "Finance"],  # text that no figure follows: rows
        ["Checked by", "Audit"],
    ]

    found = sheets.sheet_elements("NUPRC", rows, 4)

    headers = [e.text for e in found if e.kind == "table_header"]
    assert headers == [
        "CUSTOMER | VESSEL | QTY IN BBLS (BBLS)",
        "CUSTOMER | VESSEL | EXCH RATE | COMPANIES | COMPANY LIFTING",
    ]
    cells = [e.cells for e in found if e.kind == "table_row"]
    assert cells[0] == (("", "1"), ("", "2"), ("", "3"), ("", ""), ("", ""))
    assert ("QTY IN BBLS (BBLS)", "200000") in cells[1]
    assert cells[2][0] == ("CUSTOMER", "RENTAL (NAIRA)")  # under the block in force
    assert ("EXCH RATE", "401.3") in cells[3]
    assert ("COMPANY LIFTING", "0") in cells[3]  # the side table's own header
    assert cells[4][:2] == (("CUSTOMER", "Prepared by"), ("VESSEL", "Finance"))
    grid = next(e.grid for e in found if e.kind == "table")
    assert grid[1] == ("CUSTOMER", "VESSEL", "QTY IN BBLS (BBLS)", "", "")
    assert len(grid) == 8  # both header rows where they stand, and the six rows
