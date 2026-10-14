"""Spreadsheets read into elements, in the shape PDF tables take (`pdf`).

Each sheet with content becomes one table:
- its leading single-cell rows are title lines, read as headings;
- a header block is a text-only row naming at least two columns, with the text-only
  rows directly after it, when a row holding a figure follows: a multi-row header
  (a group label, the column names, a row of units such as "(MMSCF)") reads as one
  header per column, joined top to bottom;
- every other row is a table row whose cells pair each value with its column's
  header in the latest block above it, so a figure keeps its row label, column and
  unit in the text quoted. A block may come after a row of column numbers, and a
  sheet may hold several, stacked or side by side (NNPC's royalty, tax and
  production sheets; sheets-v1).

Merged cells repeat their value in every position they cover, as PDF tables do, so
a group label reaches each column under it. Numbers keep up to 15 significant
digits (no rounding of real values, no float noise); dates read as ISO dates.
"""

from datetime import date, datetime
from pathlib import Path
from typing import Any

from petroknowledge.elements import Element, grid_text, row_text

Rows = list[list[Any]]
OLE2 = b"\xd0\xcf\x11\xe0"  # an old .xls workbook; .xlsx files are zip packages


def text(value: Any) -> str:
    """A cell as written: dates as ISO dates, numbers without float noise."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        stamp = value.date() if value.time() == datetime.min.time() else value
        return stamp.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        return format(value, ".15g")
    return " ".join(str(value).split())


def is_figure(value: Any) -> bool:
    """A number or a date, as opposed to text."""
    return isinstance(value, int | float | date) and not isinstance(value, bool)


def filled(row: list[Any]) -> list[Any]:
    return [value for value in row if text(value)]


def is_line(row: list[Any]) -> bool:
    """A row holding one text: one filled cell, or one text merged across cells."""
    return len({text(value) for value in filled(row)}) == 1


def xlsx_sheets(path: Path) -> list[tuple[str, Rows]]:
    """Each sheet's name and rows, merged cells repeated."""
    import openpyxl

    # Opened as a stream: openpyxl refuses a path without an .xlsx name, and stored
    # files are named by their SHA-256.
    with path.open("rb") as file:
        book = openpyxl.load_workbook(file, data_only=True)
    found = []
    for sheet in book.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        for merged in sheet.merged_cells.ranges:
            value = sheet.cell(merged.min_row, merged.min_col).value
            for row in range(merged.min_row, merged.max_row + 1):
                for column in range(merged.min_col, merged.max_col + 1):
                    rows[row - 1][column - 1] = value
        found.append((sheet.title, rows))
    return found


def xls_sheets(path: Path) -> list[tuple[str, Rows]]:
    """Each sheet's name and rows from an old .xls workbook, merged cells repeated
    and date cells read as dates."""
    import xlrd

    book = xlrd.open_workbook(path, formatting_info=True)
    found = []
    for sheet in book.sheets():
        rows: Rows = []
        for index in range(sheet.nrows):
            row: list[Any] = []
            for cell in sheet.row(index):
                value: Any = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                row.append(value)
            rows.append(row)
        for top, bottom, left, right in sheet.merged_cells:
            value = rows[top][left]
            for row_index in range(top, bottom):
                for column in range(left, right):
                    rows[row_index][column] = value
        found.append((sheet.name, rows))
    return found


def names_columns(row: list[Any]) -> bool:
    """A text-only row naming at least two columns: a header block's first row (one
    text alone is a title or a section's name)."""
    texts = {text(value) for value in filled(row)}
    return not any(is_figure(value) for value in row) and len(texts) >= 2


def split_headers(rows: Rows) -> list[tuple[Rows, list[Any]]]:
    """Each row that is not a header, with the header block in force above it (none
    before the first). Text-only rows that no row with a figure follows are rows."""
    found: list[tuple[Rows, list[Any]]] = []
    block: Rows = []  # the header block in force
    pending: Rows = []  # text-only rows that may become the next block
    for row in rows:
        has_figure = any(is_figure(value) for value in row)
        if not has_figure and (pending or names_columns(row)):
            pending.append(row)
            continue
        if pending:  # text-only rows, then a row with a figure: a new block
            block, pending = pending, []
        found.append((block, row))
    found += [(block, line) for line in pending]
    return found


def joined_headers(header_rows: Rows, width: int) -> list[str]:
    """One header per column: the header rows' texts from top to bottom, each
    distinct text once (a merged group label is repeated across its columns)."""
    headers = []
    for column in range(width):
        parts: list[str] = []
        for row in header_rows:
            part = text(row[column]) if column < len(row) else ""
            if part and part not in parts:
                parts.append(part)
        headers.append(" ".join(parts))
    return headers


def sheet_elements(name: str, rows: Rows, number: int) -> list[Element]:
    """One sheet's title, its other title lines, then its table, header and rows."""
    rows = [row for row in rows if filled(row)]
    titles: list[str] = []
    while rows and is_line(rows[0]):
        titles.append(text(filled(rows.pop(0))[0]))
    if rows and not any(is_figure(value) for row in rows for value in row):
        split = [(rows[:1], row) for row in rows[1:]]  # text only: row 1 is the header
    else:
        split = split_headers(rows)
    found: list[Element] = []
    for index, line in enumerate(titles):
        kind = "heading" if index == 0 else "paragraph"
        found.append(
            Element(number, kind, line, (name, *titles[:1]) if index else (name,))
        )
    if not rows:
        return found
    section = (name, *titles)
    caption = " — ".join(section)
    width = max(len(row) for row in rows)
    blocks: list[list[str]] = []  # each block's joined headers, in order
    entries: list[tuple[int, list[str]]] = []  # each row's block and values
    for block, row in split:
        headers = joined_headers(block, width)
        if block and (not blocks or headers != blocks[-1]):
            blocks.append(headers)
        values = [text(row[c]) if c < len(row) else "" for c in range(width)]
        entries.append((len(blocks) - 1, values))
    # Columns empty in every header and row (a blank margin) add nothing.
    kept = [
        c
        for c in range(width)
        if any(headers[c] for headers in blocks) or any(v[c] for _, v in entries)
    ]
    blocks = [[headers[c] for c in kept] for headers in blocks]
    empty = [""] * len(kept)
    grid_rows: list[list[str]] = []
    shown = -1
    for block_index, values in entries:
        if block_index != shown and block_index >= 0:
            grid_rows.append(blocks[block_index])
            shown = block_index
        grid_rows.append([values[c] for c in kept])
    grid = tuple(tuple(row) for row in grid_rows)
    whole = grid_text(grid, caption)
    found.append(Element(number, "table", whole, section, table=number, grid=grid))
    for headers in blocks:
        if any(headers):
            line = " | ".join(header for header in headers if header)
            found.append(Element(number, "table_header", line, section, table=number))
    for index, (block_index, values) in enumerate(entries):
        headers = blocks[block_index] if block_index >= 0 else empty
        cells = tuple(zip(headers, [values[c] for c in kept], strict=True))
        line = row_text(cells, caption)
        found.append(
            Element(
                number, "table_row", line, section, table=number, row=index, cells=cells
            )
        )
    return found


def read_book(path: Path) -> list[Element]:
    """Every sheet with content, numbered from 1 in workbook order."""
    with path.open("rb") as file:
        old = file.read(4) == OLE2
    books = xls_sheets(path) if old else xlsx_sheets(path)
    found: list[Element] = []
    for number, (name, rows) in enumerate(books, 1):
        found += sheet_elements(name, rows, number)
    return found
